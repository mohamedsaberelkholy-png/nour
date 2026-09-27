import json, os, sys

request = json.load(sys.stdin)
path = request['path']
ext = os.path.splitext(path)[1].lower()
parts = []
if ext == '.pdf':
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            parts.append({'location': f'page {number}', 'text': page.extract_text() or ''})
elif ext == '.docx':
    from docx import Document
    doc = Document(path)
    text = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        text.extend(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows)
    parts.append({'location': 'document', 'text': '\n'.join(text)})
elif ext == '.xlsx':
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values: rows.append(' | '.join(values))
        parts.append({'location': f'sheet {sheet.title}', 'text': '\n'.join(rows)})
elif ext == '.pptx':
    from pptx import Presentation
    deck = Presentation(path)
    for number, slide in enumerate(deck.slides, 1):
        text = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip(): text.append(shape.text.strip())
        parts.append({'location': f'slide {number}', 'text': '\n'.join(text)})
else:
    raise ValueError('Unsupported document format')

combined = '\n\n'.join(f"[{part['location']}]\n{part['text']}" for part in parts)
print(json.dumps({'format': ext, 'parts': parts, 'text': combined[:50000]}))
